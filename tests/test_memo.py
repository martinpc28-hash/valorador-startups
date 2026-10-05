"""Memo de inversión: contenido, riesgos automáticos y exportación a PDF, Word y Excel."""

import io
import re

import numpy as np
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from src import memo as memo_mod
from src.memo_charts import is_png
from src.paths import ROOT


def _ctx(**kw) -> memo_mod.MemoContext:
    base = dict(
        company="Ejemplo SL", recommendation="Invertir", thesis="Mercado grande y equipo fuerte.",
        description="Software para pymes.", use_of_funds="Contratar ventas.", risks=["Riesgo A"], next_steps=["Paso 1"],
        currency="EUR", industry="Software (System & Application)", sector="Tecnología", stage="Serie A",
        revenue=2e6, growth=0.8, current_margin=-0.4, target_margin=0.3, margin_year=7, burn=150_000, cash=1e6, runway=6.7,
        investment=3e6, pre_money=12e6, post_money=15e6, stake=0.2, future_dilution=0.45, survival_prob=0.45, target_irr=0.5,
        verdict_title="Dentro del rango de valoraciones", lo_mid=10e6, hi_mid=20e6,
        ff_rows=[{"method": "DCF", "low": 5e6, "mid": 11e6, "high": 18e6, "range_label": "Escenarios"},
                 {"method": "Método VC", "low": 8e6, "mid": 19e6, "high": 30e6, "range_label": "Escenarios"}],
        dcf_equity=11e6, dcf_operating=20e6, dcf_pv_fcff=-3e6, dcf_pv_terminal=23e6, dcf_survival_value=10e6, cost_of_equity=0.27, cost_of_capital=0.27,
        mature_coc=0.094, beta_used=4.65, beta_type="total", stable_growth=0.03,
        projection=pd.DataFrame({"Año": range(1, 11), "Ingresos": np.linspace(3e6, 40e6, 10),
                                 "Margen operativo": np.linspace(-0.3, 0.3, 10), "FCFF": np.linspace(-2e6, 6e6, 10)}),
        top_sensitivity="Crecimiento ±30 %", exit_year=6, exit_multiple=11.4, small_cap_multiple=1.23, exit_value=500e6,
        vc_rate=0.5, vc_mode="IRR objetivo de la etapa", vc_pre_money=19e6, vc_post_money=22e6, required_stake=0.14,
        moic=19.5, irr=0.64, expected_moic=8.8, stake_exit=0.11, proceeds=58e6,
        multiples=pd.DataFrame({"Método": ["EV/Sales"], "Múltiplo": ["11,41x"], "Valor del equity": ["€ 16 M"], "Fuente": ["damodaran"]}),
        illiquidity_discount=0.3, scenarios=pd.DataFrame({"Escenario": ["Base"], "DCF": ["€ 11 M"]}),
        scenarios_raw=pd.DataFrame({"Escenario": ["Pesimista", "Base", "Optimista"], "DCF": [0.0, 11e6, 30e6],
                                    "Método VC (pre-money)": [6e6, 19e6, 45e6], "Múltiplos (EV/Sales)": [8e6, 16e6, 25e6]}),
        cash_projection=pd.DataFrame({"Mes": range(1, 49), "Caja": np.linspace(4e6, -2e6, 48),
                                      "Consumo de caja": np.linspace(150e3, -10e3, 48)}),
        mc_prob_fail=0.55, mc_prob_target=0.43, mc_prob_loss=0.55, mc_moic_mean=11.7, mc_moic_pct={"P10": 5.5, "P50": 18.3, "P90": 55.0},
        mc_moic=np.r_[np.zeros(50), np.linspace(1, 40, 50)], mc_survived=np.r_[np.zeros(50, bool), np.ones(50, bool)],
        capital_need=25e6, funding_gap=21e6, implied_dilution=0.58,
        assumptions=pd.DataFrame([("Crecimiento", "80 %", "usuario")], columns=["Supuesto", "Valor", "Origen"]),
        provenance=pd.DataFrame([("Beta", "1,25", "damodaran", "Software", "2026-01-09")],
                                columns=["Dato", "Valor", "Fuente", "Industria usada", "Fecha"]),
    )
    base.update(kw)
    return memo_mod.MemoContext(**base)


def test_auto_risks_detect_the_signals():
    risks = " ".join(memo_mod.auto_risks(_ctx()))
    assert "Runway de 6,7 meses" in risks
    assert "faltan" in risks and "dilución" in risks
    assert "fracaso del 55 %" in risks
    assert "más del doble" in risks  # 11,4x frente a 1,23x de las cotizadas pequeñas
    assert "ilustrativos" in risks
    calm = memo_mod.auto_risks(_ctx(runway=30, funding_gap=0, mc_prob_fail=0.1, exit_multiple=2.0, growth=0.3,
                                    current_margin=0.1, mc_prob_target=0.6, top_sensitivity=""))
    assert len(calm) == 1  # solo el aviso de supuestos ilustrativos


def test_memo_structure_and_no_long_dashes():
    memo = memo_mod.build_memo(_ctx())
    titles = [s.title for s in memo.sections]
    assert titles[0] == "1. Resumen ejecutivo"
    assert any("Riesgos y mitigantes" in t for t in titles) and titles[-1].startswith("Anexo")
    text = memo_mod.memo_text(memo)
    assert "Recomendación: Invertir" in text and "Mercado grande" in text
    assert not re.search("[—–]", text), "sin guiones largos"
    figs = [b for s in memo.sections for b in s.blocks if isinstance(b, memo_mod.Figure)]
    assert len(figs) == 6 and all(is_png(f.png) for f in figs)  # métodos, proyección, puente DCF, MOIC, escenarios, caja
    assert "Analista" not in memo.meta and "Fondo" not in memo.meta  # formato estándar, sin firma


def test_exports_open_and_contain_the_numbers():
    ctx = _ctx()
    memo = memo_mod.build_memo(ctx)
    pdf = memo_mod.to_pdf(memo)
    assert pdf[:5] == b"%PDF-" and len(pdf) > 20_000

    from docx import Document
    doc = Document(io.BytesIO(memo_mod.to_docx(memo)))
    body = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Memo de inversión: Ejemplo SL" in body and "Resumen ejecutivo" in body and "€ 12,00 M" in body
    assert len(doc.inline_shapes) == 6
    assert "Lectura analítica" in body and "Puntos clave" in body

    xl = pd.read_excel(io.BytesIO(memo_mod.to_xlsx(ctx, memo)), sheet_name=None)
    assert {"Resumen", "Métodos", "Proyección DCF", "Supuestos", "Fuentes", "Riesgos", "Lectura analítica"} <= set(xl)
    resumen = xl["Resumen"].set_index("Concepto")["Valor"]
    assert float(resumen["Pre-money"]) == 12e6


def test_app_generates_memo_end_to_end():
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    assert not any(t.label in ("Analista", "Fondo") for t in at.text_input)  # memo estándar, sin firma
    next(t for t in at.text_area if t.label == "Tesis de inversión").input("Mercado en crecimiento.")
    at.run()
    risks = next(t for t in at.text_area if t.label.startswith("Riesgos y mitigantes")).value
    assert "Runway" in risks or "ilustrativos" in risks  # prellenado con los riesgos detectados
    next(b for b in at.button if b.label == "📝 Generar memo").click()
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    files = at.session_state["memo_files"]
    assert files["pdf"][:5] == b"%PDF-"
    text = memo_mod.memo_text(files["obj"])
    assert "Puntos clave" in text and "Lectura analítica" in text and "Mercado en crecimiento." in text
    assert not re.search("[—–]", text)
    assert sum(isinstance(b, memo_mod.Figure) for s in files["obj"].sections for b in s.blocks) == 6


def test_analytical_reading_cites_the_numbers():
    from src import memo_analysis as ma

    c = _ctx()
    kp = " ".join(ma.key_points(c))
    assert "Precio dentro del rango" in kp and "€ 12,00 M" in kp  # 12 M entre DCF (11 M) y método VC (19 M)
    assert "MOIC de 19,50x supera" in kp  # frente a 1,5^6 = 11,39x
    assert "Riesgo:" in kp and "Caja:" in kp
    assert any("Regla del 40" in t and "40 %" in t for t in ma.read_company(c))  # 80 % + (-40 %) = 40 %
    assert any("115 %" in t for t in ma.read_dcf(c))  # 23 M de terminal sobre 20 M de valor operativo
    assert any("ningún método" in t for t in ma.read_scenarios(c))  # el pesimista no sostiene 12 M
    assert any("se agota en el mes" in t for t in ma.read_cash(c))
    expensive = " ".join(ma.key_points(_ctx(pre_money=40e6)))
    assert "Precio exigente" in expensive


def test_analytical_reading_skips_missing_data():
    from src import memo_analysis as ma

    c = _ctx(projection=None, scenarios_raw=None, cash_projection=None, mc_moic_pct={}, ff_rows=[])
    assert ma.read_scenarios(c) == [] and ma.read_valuation(c) == []
    memo_mod.build_memo(c)  # sin errores aunque falten datos
